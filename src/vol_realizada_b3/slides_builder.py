"""Construcao da apresentacao PowerPoint final."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .config import PROJECT_ROOT, load_config


NAVY = RGBColor(24, 54, 93)
BLUE = RGBColor(55, 105, 150)
RED = RGBColor(180, 55, 55)
DARK = RGBColor(40, 44, 52)
LIGHT = RGBColor(245, 247, 250)
GRAY = RGBColor(105, 112, 122)


def _pct(value: float) -> str:
    return "n/d" if not np.isfinite(value) else f"{100 * value:.1f}%"


class AcademicDeck:
    """Helper pequeno para um layout academico consistente."""

    def __init__(self) -> None:
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(13.333)
        self.presentation.slide_height = Inches(7.5)

    def _background(self, slide: Any) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = LIGHT
        band = slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(13.333), Inches(0.16)
        )
        band.fill.solid()
        band.fill.fore_color.rgb = NAVY
        band.line.fill.background()

    def _footer(self, slide: Any, number: int) -> None:
        box = slide.shapes.add_textbox(
            Inches(0.55), Inches(7.08), Inches(12.2), Inches(0.22)
        )
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = f"FGV EESP | Volatilidade Realizada B3 | {number}"
        paragraph.font.size = Pt(8)
        paragraph.font.color.rgb = GRAY
        paragraph.alignment = PP_ALIGN.RIGHT

    @staticmethod
    def _add_picture_fit(
        slide: Any,
        image: Path,
        left: int,
        top: int,
        max_width: int,
        max_height: int,
    ) -> None:
        with Image.open(image) as source:
            width_px, height_px = source.size
        ratio = width_px / height_px
        box_ratio = max_width / max_height
        if ratio >= box_ratio:
            width = max_width
            height = int(max_width / ratio)
        else:
            height = max_height
            width = int(max_height * ratio)
        x = left + (max_width - width) // 2
        y = top + (max_height - height) // 2
        slide.shapes.add_picture(str(image), x, y, width=width, height=height)

    def title_slide(self, title: str, subtitle: str) -> None:
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[6]
        )
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = NAVY
        title_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.65), Inches(11.7), Inches(2.2)
        )
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.font.size = Pt(30)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = RGBColor(255, 255, 255)
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.85), Inches(4.35), Inches(11.5), Inches(1.0)
        )
        subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
        subtitle_paragraph.text = subtitle
        subtitle_paragraph.font.size = Pt(18)
        subtitle_paragraph.font.color.rgb = RGBColor(220, 229, 240)

    def content_slide(
        self,
        title: str,
        bullets: Iterable[str] = (),
        image: Path | None = None,
        note: str | None = None,
    ) -> None:
        slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[6]
        )
        self._background(slide)
        number = len(self.presentation.slides)
        self._footer(slide, number)

        title_box = slide.shapes.add_textbox(
            Inches(0.55), Inches(0.35), Inches(12.1), Inches(0.55)
        )
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.font.size = Pt(23)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = NAVY

        bullet_list = list(bullets)
        has_image = image is not None and image.exists()
        text_width = Inches(4.5) if has_image else Inches(11.9)
        body = slide.shapes.add_textbox(
            Inches(0.65), Inches(1.2), text_width, Inches(5.45)
        )
        text_frame = body.text_frame
        text_frame.word_wrap = True
        text_frame.clear()
        for index, item in enumerate(bullet_list):
            paragraph = (
                text_frame.paragraphs[0]
                if index == 0
                else text_frame.add_paragraph()
            )
            paragraph.text = item
            paragraph.level = 0
            paragraph.font.size = Pt(16)
            paragraph.font.color.rgb = DARK
            paragraph.space_after = Pt(9)
            paragraph.text = f"• {paragraph.text}"

        if has_image:
            self._add_picture_fit(
                slide,
                image,
                left=Inches(5.35),
                top=Inches(1.15),
                max_width=Inches(7.45),
                max_height=Inches(5.6),
            )
        if note:
            note_box = slide.shapes.add_textbox(
                Inches(0.7), Inches(6.65), Inches(12.0), Inches(0.3)
            )
            paragraph = note_box.text_frame.paragraphs[0]
            paragraph.text = note
            paragraph.font.size = Pt(9)
            paragraph.font.italic = True
            paragraph.font.color.rgb = GRAY


def _result_context() -> dict[str, Any]:
    tables = PROJECT_ROOT / "outputs" / "tables"
    coverage = pd.read_csv(tables / "data_coverage_by_ticker.csv")
    realized = pd.read_csv(tables / "realized_measures_summary.csv")
    jumps = pd.read_csv(tables / "jump_summary.csv")
    groups = pd.read_csv(tables / "group_comparison.csv")
    garch = pd.read_csv(tables / "garch_summary.csv")
    events = pd.read_csv(tables / "event_window_summary.csv")
    daily_quality = pd.read_csv(
        PROJECT_ROOT / "data" / "interim" / "daily_data_quality.csv"
    )
    measures = pd.read_csv(
        PROJECT_ROOT / "data" / "processed" / "realized_measures.csv"
    )
    included = coverage.loc[coverage["status"].eq("included")]
    excluded = coverage.loc[coverage["status"].eq("excluded")]
    top_rvol = realized.sort_values("mean_rvol_annualized", ascending=False).iloc[0]
    top_jump = jumps.sort_values("jump_day_percentage", ascending=False).iloc[0]
    return {
        "coverage": coverage,
        "included": included,
        "excluded": excluded,
        "groups": groups,
        "garch": garch,
        "events": events,
        "daily_quality": daily_quality,
        "measures": measures,
        "top_rvol": top_rvol,
        "top_jump": top_jump,
    }


def run_slides_builder(
    config: dict[str, Any] | None = None,
) -> Path:
    """Gera os 18 slides especificados."""
    config = config or load_config()
    context = _result_context()
    figures = PROJECT_ROOT / "outputs" / "figures"
    effective_close = (
        str(context["daily_quality"]["effective_market_close"].mode().iloc[0])
        if "effective_market_close" in context["daily_quality"]
        else config["data"]["market_close"]
    )
    expected_candles = int(
        context["daily_quality"]["expected_candles"].mode().iloc[0]
    )
    groups = context["groups"]
    core = groups.loc[groups["grupo"].eq("core_liquid")].iloc[0]
    high_vol = groups.loc[
        groups["grupo"].eq("high_vol_growth_candidate")
    ].iloc[0]
    rvol_ratio = high_vol["mean_rvol_annualized"] / core["mean_rvol_annualized"]
    jump_ratio = high_vol["jump_frequency"] / core["jump_frequency"]
    top_day = context["measures"].nlargest(1, "rvol_annualized").iloc[0]
    deck = AcademicDeck()
    deck.title_slide(
        config["project"]["title"],
        f"{config['project']['student_name']} | {config['project']['institution']} | Trabalho final",
    )
    deck.content_slide(
        "Objetivo e pergunta do trabalho",
        [
            "Estimar volatilidade realizada com retornos intradiarios de 5 minutos.",
            "Separar variacao continua e jumps com BV, JV e teste estatistico.",
            "Comparar ativos liquidos com candidatas growth/high-vol/small caps.",
            "Avaliar implicacoes para monitoramento e gestao de risco.",
        ],
    )
    deck.content_slide(
        "Dados",
        [
            "Fonte: Yahoo Finance via yfinance.",
            "Frequencia: candles de 5 minutos; janela solicitada: 60 dias.",
            f"Ativos incluidos: {len(context['included'])} de {len(context['coverage'])} candidatos.",
            "Limitacao: historico intradiario curto e ausencia de dados tick-by-tick.",
        ],
        figures / "data_coverage_by_ticker.png",
    )
    deck.content_slide(
        "Estrategia de amostra",
        [
            "Core: grandes acoes liquidas para comparabilidade.",
            "Complementar: growth, small caps e ativos potencialmente mais volateis.",
            "Criterios: 30 dias validos, cobertura >= 70%, >= 40 retornos/dia.",
            "Precos positivos, volume relevante e precos parados <= 50%.",
        ],
    )
    deck.content_slide(
        "Tratamento dos dados",
        [
            f"Timezone America/Sao_Paulo; teto 17:55 e fechamento efetivo {effective_close}.",
            f"Grade regular de 5 minutos com {expected_candles} candles sustentados pela fonte.",
            "Remocao de duplicatas e precos invalidos.",
            "Forward-fill somente dentro do dia; primeiro retorno diario removido.",
            "Fechamento inferido evita uma cauda artificial de retornos zero.",
        ],
    )
    deck.content_slide(
        "Metodologia",
        [
            "RV = soma dos retornos intradiarios ao quadrado.",
            "RVol diaria = raiz de RV; anualizada = raiz de 252 x RV.",
            "BV aproxima variacao continua; JV = max(RV - BV, 0).",
            "Teste BNS com tripower quarticity e quantil normal de 99%.",
            "GARCH(1,1) em retornos diarios close-to-close.",
        ],
    )
    deck.content_slide(
        "Cobertura da amostra",
        [
            f"Incluidos: {', '.join(context['included']['ticker']) or 'nenhum'}.",
            f"Excluidos: {', '.join(context['excluded']['ticker']) or 'nenhum'}.",
            "Exclusoes evitam que baixa liquidez contamine jumps e comparacoes.",
        ],
        figures / "data_coverage_by_ticker.png",
    )
    deck.content_slide(
        "Volatilidade realizada ao longo do tempo",
        [
            f"Maior ticker-dia: {top_day['ticker']} com {_pct(top_day['rvol_annualized'])}.",
            "O maior pico nao foi jump: RV e BV subiram juntas.",
            "Persistencia visual sugere clustering de volatilidade.",
            "Growth/high-vol opera em patamar superior ao core na maior parte da janela.",
        ],
        figures / "realized_volatility_time_series.png",
    )
    deck.content_slide(
        "Comparacao entre ativos",
        [
            f"Maior RVol anualizada media: {context['top_rvol']['ticker']} "
            f"({_pct(context['top_rvol']['mean_rvol_annualized'])}).",
            "CASH3, LWSA3 e MGLU3 ocupam as tres primeiras posicoes.",
            "TOTS3 e o ativo core de maior volatilidade media.",
            "Boxplots comparam nivel, dispersao e caudas entre ativos.",
        ],
        figures / "rvol_boxplot_by_ticker.png",
    )
    group_bullets = [
        f"{row.grupo}: RVol={_pct(row.mean_rvol_annualized)}, "
        f"jumps={_pct(row.jump_frequency)}."
        for row in groups.itertuples()
    ]
    group_bullets.extend(
        [
            f"Razao de RVol complementar/core: {rvol_ratio:.2f}x.",
            f"Razao da frequencia de jumps: {jump_ratio:.2f}x.",
            "Resultado condicionado aos filtros: apenas 3 candidatas passaram.",
        ]
    )
    deck.content_slide(
        "Core liquido versus growth/high-vol",
        group_bullets,
        figures / "core_vs_high_vol_comparison.png",
    )
    deck.content_slide(
        "Bipower Variation e jumps",
        [
            "BV usa produtos de retornos absolutos adjacentes.",
            "RV acima de BV gera JV positiva.",
            "Diferenca economica: risco continuo versus movimentos descontínuos.",
            "JV positiva nao basta: o teste usa TQ e valor critico de 99%.",
        ],
        figures / "rv_vs_bv.png",
    )
    deck.content_slide(
        "Frequencia e intensidade dos jumps",
        [
            f"Maior frequencia: {context['top_jump']['ticker']} "
            f"({_pct(context['top_jump']['jump_day_percentage'])}).",
            "LWSA3 ficou em segundo; VALE3 e ITUB4 tiveram as menores frequencias.",
            "Jump share mede a parcela estimada de RV associada a saltos.",
            "Jumps afetam risco de gap, execucao e stops entre observacoes.",
        ],
        figures / "jump_frequency_by_ticker.png",
    )
    successful = context["garch"].loc[context["garch"]["fit_status"].eq("ok")]
    persistence = (
        f"Persistencia mediana alpha+beta: {successful['alpha_plus_beta'].median():.3f}."
        if not successful.empty
        else "Ajustes GARCH limitados pela amostra curta."
    )
    deck.content_slide(
        "GARCH versus volatilidade realizada",
        [
            persistence,
            "GARCH captura persistencia com resposta suavizada.",
            "RVol reage diretamente a picos intradiarios e jumps.",
            "Maior correlacao GARCH-RVol ocorreu em B3SA3, ainda moderada.",
            "Close-to-close inclui overnight; RVol intradiaria nao.",
            "Amostra de 59 retornos: parametros devem ser lidos com cautela.",
        ],
        figures / "garch_vs_realized_all.png",
    )
    valid_events = context["events"].loc[
        context["events"].get("status", pd.Series(dtype=str)).eq("ok")
    ]
    deck.content_slide(
        "Eventos de resultados",
        [
            (
                f"{len(valid_events)} eventos de earnings analisados em t-1, t e t+1."
                if not valid_events.empty
                else "Nenhum evento manual verificavel foi preenchido."
            ),
            (
                f"RVol da janela: "
                f"{valid_events['event_vs_normal_ratio'].mean() - 1:.1%} acima "
                "dos dias normais."
                if not valid_events.empty
                else "Comparacao de RVol indisponivel."
            ),
            (
                f"Jump days: "
                f"{valid_events['event_window_jump_frequency'].mean():.1%} nas "
                f"janelas vs {valid_events['normal_jump_frequency'].mean():.1%} "
                "nos dias normais."
                if not valid_events.empty
                else "Comparacao de jumps indisponivel."
            ),
            "Maior elevacao relativa em ABEV3; PETR4 ficou abaixo da media normal.",
            "Evidencia descritiva; fonte terceirizada e sem identificacao causal.",
        ],
        figures / "event_window_volatility.png",
    )
    deck.content_slide(
        "Implicacoes para gestao de risco",
        [
            "VaR e expected shortfall podem subestimar perdas em jump days.",
            "RVol intradiaria melhora o monitoramento de mudancas rapidas.",
            "Jump risk maior pede limites, margens e sizing conservadores.",
            "Comovimento da volatilidade reduz diversificacao em estresse.",
            "Liquidez baixa exige cautela com falsos jumps.",
        ],
    )
    deck.content_slide(
        "Limitacoes",
        [
            "Historico intradiario limitado pelo Yahoo Finance.",
            "Ruido de microestrutura e ausencia de ticks.",
            "Resultados dependem da frequencia e sincronizacao.",
            "Small caps podem ter candles faltantes e precos parados.",
            "GARCH diario estimado em amostra curta.",
        ],
    )
    deck.content_slide(
        "Conclusao",
        [
            f"O ranking observado aponta {context['top_rvol']['ticker']} como maior RVol media.",
            "RV, BV, jumps e GARCH oferecem dimensoes complementares do risco.",
            "Qualidade intradiaria deve ser avaliada antes de inferir jump risk.",
            "Extensoes: ticks B3, HAR-RV, realized kernels e eventos verificados.",
        ],
        figures / "asset_risk_ranking.png",
    )
    deck.content_slide(
        "Referencias",
        [
            "Andersen, Bollerslev, Diebold e Labys (2003).",
            "Barndorff-Nielsen e Shephard (2004, 2006).",
            "Bollerslev (1986).",
            "Documentacao do pacote arch.",
            "Documentacao do yfinance.",
        ],
        note="Referencias completas em report/referencias.md.",
    )

    output = (
        PROJECT_ROOT
        / "outputs"
        / "slides"
        / "trabalho_volatilidade_realizada_b3.pptx"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.presentation.save(output)
    return output
